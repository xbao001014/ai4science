"""Generate a short PPT summarizing updates since mid/late July 2026."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "output" / "updates_since_20260723.pptx"

BG = RGBColor(0x0F, 0x1C, 0x24)
PANEL = RGBColor(0x16, 0x2A, 0x36)
ACCENT = RGBColor(0x2A, 0x9D, 0x8F)
ACCENT2 = RGBColor(0xE9, 0xC4, 0x6A)
TEXT = RGBColor(0xF4, 0xF7, 0xF5)
MUTED = RGBColor(0xA8, 0xB8, 0xC0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WARN = RGBColor(0xE0, 0x7A, 0x5F)


def _run(run, size=18, bold=False, color=TEXT, font="Microsoft YaHei"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, prs):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    sp_tree = slide.shapes._spTree
    sp = shape._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)


def add_bar(slide, prs):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.08)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=18,
    bold=False,
    color=TEXT,
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _run(run, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, left, top, width, height, items, size=16, color=TEXT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = "•  " + item
        _run(run, size=size, color=color)
    return box


def card(slide, left, top, width, height, title, body_lines, title_color=ACCENT):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = PANEL
    shp.line.fill.background()
    add_text(
        slide,
        left + Inches(0.25),
        top + Inches(0.2),
        width - Inches(0.4),
        Inches(0.4),
        title,
        size=16,
        bold=True,
        color=title_color,
    )
    add_bullets(
        slide,
        left + Inches(0.2),
        top + Inches(0.65),
        width - Inches(0.35),
        height - Inches(0.8),
        body_lines,
        size=13,
        color=MUTED,
    )


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 Title
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_bar(s, prs)
    add_text(
        s,
        Inches(0.8),
        Inches(2.0),
        Inches(11.5),
        Inches(0.5),
        "病理 AI 文献知识图谱工作流",
        size=18,
        color=ACCENT,
    )
    add_text(
        s,
        Inches(0.8),
        Inches(2.6),
        Inches(11.5),
        Inches(1.0),
        "系统更新汇报",
        size=40,
        bold=True,
        color=WHITE,
    )
    add_text(
        s,
        Inches(0.8),
        Inches(3.9),
        Inches(11.5),
        Inches(0.5),
        "相对 2026-07-23 时点（仓库基线 7/19 tip）→ 当前主线",
        size=18,
        color=MUTED,
    )
    add_text(
        s,
        Inches(0.8),
        Inches(5.4),
        Inches(11.5),
        Inches(0.4),
        "抽取治理  ·  研究类型策略  ·  改进建议  ·  可迁移热点  ·  抽取 Demo",
        size=16,
        color=ACCENT2,
    )
    add_text(
        s, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.35), "2026-07-29", size=14, color=MUTED
    )

    # 2 Overview
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_bar(s, prs)
    add_text(
        s,
        Inches(0.7),
        Inches(0.35),
        Inches(12),
        Inches(0.5),
        "更新总览：五条主线",
        size=28,
        bold=True,
        color=WHITE,
    )
    add_text(
        s,
        Inches(0.7),
        Inches(0.95),
        Inches(12),
        Inches(0.35),
        "约 59 次提交 · 109 文件 · 从「可视化 + 基础热点」推进到可审计抽取与可解释选题信号",
        size=14,
        color=MUTED,
    )
    overview = [
        (
            "01  抽取治理",
            [
                "全文 Pass2 reconcile",
                "关系 status / 合并",
                "文献平台 Dataset 过滤",
                "Pilot 按 PMID 重跑",
            ],
        ),
        (
            "02  Study-type",
            [
                "研究类型分类对齐",
                "Prompt packs + 策略矩阵",
                "Survey / cover 边",
                "分层 QA 统计",
            ],
        ),
        (
            "03  改进建议",
            [
                "Pass2 推荐解析落库",
                "主题 × 动作分桶",
                "UI / Agent 工具",
                "局限联动建议",
            ],
        ),
        (
            "04  可迁移热点",
            [
                "去掉笛卡尔积假空白",
                "合格 Task 桥接",
                "抽取期 Task 质控",
                "审计 CLI",
            ],
        ),
        (
            "05  抽取 Demo",
            [
                "全文↔抽取双栏 HTML",
                "证据定位增强",
                "离线可演示导出",
                "Pilot 对照脚本",
            ],
        ),
    ]
    for i, (title, lines) in enumerate(overview):
        x = Inches(0.45) + Inches(i * 2.55)
        card(s, x, Inches(1.6), Inches(2.4), Inches(4.8), title, lines)

    # 3 Reconcile
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_bar(s, prs)
    add_text(
        s,
        Inches(0.7),
        Inches(0.35),
        Inches(12),
        Inches(0.5),
        "抽取治理：Pass2 Reconcile",
        size=28,
        bold=True,
        color=WHITE,
    )
    add_text(
        s,
        Inches(0.7),
        Inches(1.0),
        Inches(12),
        Inches(0.4),
        "目标：压低平台噪声、稳住合并，只保留研究真实使用的数据集与有效边",
        size=15,
        color=MUTED,
    )
    card(
        s,
        Inches(0.6),
        Inches(1.7),
        Inches(5.9),
        Inches(4.6),
        "做了什么",
        [
            "关系 status / superseded；读侧忽略失效边",
            "全文级 reconcile：合并、纠错、平台黑名单",
            "postprocess 丢弃 PubMed 等文献平台 Dataset",
            "CLI：按 PMID 列表 extract / reconcile",
        ],
    )
    card(
        s,
        Inches(6.8),
        Inches(1.7),
        Inches(5.9),
        Inches(4.6),
        "价值",
        [
            "配置开关与输入长度上限可控",
            "sidecar / binding 结构支撑质检回放",
            "降低 Pass2 过删与错误合并风险",
            "为全库重抽提供可验收基线",
        ],
        ACCENT2,
    )

    # 4 Study-type
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_bar(s, prs)
    add_text(
        s,
        Inches(0.7),
        Inches(0.35),
        Inches(12),
        Inches(0.5),
        "Study-type：按论文类型约束抽取",
        size=28,
        bold=True,
        color=WHITE,
    )
    add_text(
        s,
        Inches(0.7),
        Inches(1.0),
        Inches(12),
        Inches(0.4),
        "算法 / 综述 / 数据集基准等使用不同 prompt 与策略矩阵",
        size=15,
        color=MUTED,
    )
    add_bullets(
        s,
        Inches(0.9),
        Inches(1.7),
        Inches(11.5),
        Inches(5),
        [
            "分类器与学者视角对齐；分研究类型的 section prompt packs",
            "Policy matrix：postprocess 中 deny / remap / dataset_mode",
            "综述类：survey / cover 边；抑制不当 APPLIES_METHOD 与数据集噪声",
            "Pass2 写入受 enable_new 门控，避免类型越权落库",
            "配套分层 pilot QA 与关系类型统计，便于抽后验收",
        ],
        size=17,
    )

    # 5 Improvement
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_bar(s, prs)
    add_text(
        s,
        Inches(0.7),
        Inches(0.35),
        Inches(12),
        Inches(0.5),
        "改进建议：从局限到可执行动作",
        size=28,
        bold=True,
        color=WHITE,
    )
    add_text(
        s,
        Inches(0.7),
        Inches(1.0),
        Inches(12),
        Inches(0.4),
        "作者局限 → 结构化推荐 → 主题分桶 → UI / 辩论工具",
        size=15,
        color=MUTED,
    )
    flow = [
        ("Pass2", "解析 improvement\nrecommendations"),
        ("DB", "论文级 sidecar\n持久化"),
        ("分析", "主题 × 动作\n分桶聚合"),
        ("产品", "Gap UI 展示\nIdea / Agent 工具"),
    ]
    for i, (h, b) in enumerate(flow):
        x = Inches(0.7) + Inches(i * 3.1)
        shp = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(2.8), Inches(2.4)
        )
        shp.fill.solid()
        shp.fill.fore_color.rgb = PANEL
        shp.line.fill.background()
        add_text(
            s,
            x + Inches(0.15),
            Inches(2.25),
            Inches(2.5),
            Inches(0.4),
            h,
            size=18,
            bold=True,
            color=ACCENT,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            s,
            x + Inches(0.15),
            Inches(2.85),
            Inches(2.5),
            Inches(1.3),
            b,
            size=14,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )
        if i < 3:
            add_text(
                s,
                x + Inches(2.75),
                Inches(2.9),
                Inches(0.4),
                Inches(0.4),
                "→",
                size=22,
                bold=True,
                color=ACCENT2,
                align=PP_ALIGN.CENTER,
            )
    add_text(
        s,
        Inches(0.9),
        Inches(5.0),
        Inches(11.5),
        Inches(1.4),
        "区分 synthesized（轻度综合）与 author_stated（贴近作者原述），\n减少「只有局限、没有下一步」的空白。",
        size=16,
        color=TEXT,
    )

    # 6 Transferable
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_bar(s, prs)
    add_text(
        s,
        Inches(0.7),
        Inches(0.35),
        Inches(12),
        Inches(0.5),
        "每周热点：假空白 → 可迁移候选",
        size=28,
        bold=True,
        color=WHITE,
    )
    card(
        s,
        Inches(0.55),
        Inches(1.3),
        Inches(6.0),
        Inches(5.2),
        "旧逻辑问题",
        [
            "热方法 Top-N × 热疾病 Top-N 笛卡尔积",
            "方法 A 热于病种 A、方法 B 热于病种 B",
            "并不意味「方法 A × 病种 B」值得做",
            "「没共现」≠「真空白」",
        ],
        WARN,
    )
    card(
        s,
        Inches(6.8),
        Inches(1.3),
        Inches(6.0),
        Inches(5.2),
        "新逻辑",
        [
            "方法升温 + 目标病相关 + 共现稀疏",
            "方法已在其他病种用过（可迁移先验）",
            "存在合格 Task（ok）桥接证据",
            "输出 bridge_task / mode / support_diseases",
            "抽取侧：泛化 Task 丢弃 + synonym",
            "CLI：task-quality-audit 只读审计",
        ],
        ACCENT,
    )

    # 7 Demo
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_bar(s, prs)
    add_text(
        s,
        Inches(0.7),
        Inches(0.35),
        Inches(12),
        Inches(0.5),
        "抽取 Demo：可对外讲清楚的对照页",
        size=28,
        bold=True,
        color=WHITE,
    )
    add_bullets(
        s,
        Inches(0.9),
        Inches(1.4),
        Inches(11.5),
        Inches(5.5),
        [
            "离线单页 HTML：左全文 / 右抽取要素，点击证据高亮",
            "从 SQLite 装载 active 关系；缺 PMID / 无章节 / 无抽取则 fail-fast",
            "证据匹配增强：分号拼接、省略号片段、Unicode 破折号 / NBSP",
            "Pass2 的 fulltext_reconcile 可回退到 discussion 等叙事段搜索",
            "配套 pilot 选 PMID / 快照 / 对比脚本，支撑重抽前后对照",
        ],
        size=17,
    )

    # 8 Next
    s = prs.slides.add_slide(blank)
    add_bg(s, prs)
    add_bar(s, prs)
    add_text(
        s,
        Inches(0.7),
        Inches(0.35),
        Inches(12),
        Inches(0.5),
        "下一步建议",
        size=28,
        bold=True,
        color=WHITE,
    )
    nexts = [
        ("全库重抽", ["让 Task synonym / reject、study-type 策略真正落库"]),
        ("跑审计", ["main.py task-quality-audit，查看 Task 分层与近重复"]),
        ("热点周更", ["坚持 hotspot-report 快照，启用 WoW + 可迁移列表"]),
        ("选题闭环", ["可迁移候选 → Gap 辩论 → 数据可行性 → 方案"]),
    ]
    for i, (h, lines) in enumerate(nexts):
        x = Inches(0.55) + Inches((i % 2) * 6.3)
        y = Inches(1.4) + Inches((i // 2) * 2.5)
        card(s, x, y, Inches(6.0), Inches(2.2), h, lines, ACCENT2 if i % 2 else ACCENT)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"wrote {path} ({path.stat().st_size} bytes)")
